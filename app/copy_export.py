"""Export generated copy for Walmart content upload — Excel (.xlsx) and CSV.

Both formats carry the same columns (one row per completed item), mapping to the
content fields a user fills on a Walmart bulk upload: Item ID, Product URL,
Product Name, Site Description, and each Key Feature as its own column. The Excel
path uses openpyxl (pure-Python, no system libraries, like the PDF export); the
CSV path uses the stdlib so users without Excel still have a universal option.

Both take the same per-item view dicts the results page uses (see
``pages._copy_row_view``): a completed item has ``status == 'done'`` and a ``new``
dict of ``{title, bullets[], description}``.
"""

import csv
import io
import logging

logger = logging.getLogger(__name__)

# Cap the per-bullet columns so a single outlier item can't explode the sheet
# width; bullets beyond this are dropped from the export.
_MAX_BULLET_COLS = 10

_BASE_HEADERS = ["Item ID", "Product URL", "Product Name", "Site Description"]


def _done_items(items: list[dict]) -> list[dict]:
    """Only completed items carry generated copy worth exporting."""
    return [it for it in items if it.get("status") == "done" and it.get("new")]


def _layout(done: list[dict]) -> tuple[list[str], int]:
    """Return the header row and the number of Key Feature columns to emit.

    The Key Feature column count is the most bullets any item has (capped), so the
    table is rectangular without a trailing pile of empty columns.
    """
    max_bullets = min(
        _MAX_BULLET_COLS,
        max((len(it["new"].get("bullets") or []) for it in done), default=0),
    )
    headers = _BASE_HEADERS + [f"Key Feature {i}" for i in range(1, max_bullets + 1)]
    return headers, max_bullets


def _row_values(it: dict, max_bullets: int) -> list[str]:
    """Flatten one completed item's NEW copy into the row order of the headers."""
    new = it["new"]
    bullets = (new.get("bullets") or [])[:max_bullets]
    values = [
        it.get("item_id") or "",
        it.get("url") or "",
        new.get("title") or "",
        new.get("description") or "",
        *bullets,
    ]
    # Pad short bullet lists so every row matches the header width.
    values += [""] * (len(_BASE_HEADERS) + max_bullets - len(values))
    return values


def build_copy_csv(items: list[dict]) -> bytes:
    """Return the generated copy as CSV bytes (UTF-8 with BOM).

    The BOM makes Excel/Sheets detect UTF-8 so accented characters survive. Only
    completed items are included.
    """
    done = _done_items(items)
    headers, max_bullets = _layout(done)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    for it in done:
        writer.writerow(_row_values(it, max_bullets))
    logger.info("Built copy CSV: %d row(s)", len(done))
    # utf-8-sig writes the BOM Excel looks for; newlines already handled by csv.
    return buf.getvalue().encode("utf-8-sig")


def build_copy_xlsx(items: list[dict]) -> bytes:
    """Return the generated copy as an .xlsx workbook (bytes).

    One sheet, a bold header row, sensible column widths, wrapped long-text cells,
    and a frozen header. Only completed items are included.
    """
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter

    done = _done_items(items)
    headers, max_bullets = _layout(done)

    wb = Workbook()
    ws = wb.active
    ws.title = "New Copy"

    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")

    for it in done:
        ws.append(_row_values(it, max_bullets))

    # Widths: Item ID / URL narrow, Product Name medium, Site Description wide,
    # each Key Feature medium. Columns are 1-based: A=1 … D=4, Key Feature i = 4+i.
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 42
    ws.column_dimensions["C"].width = 50
    ws.column_dimensions["D"].width = 80
    for i in range(1, max_bullets + 1):
        ws.column_dimensions[get_column_letter(4 + i)].width = 48

    # Wrap the long-text columns (Product Name, Site Description, Key Features).
    wrap_cols = {3, 4} | {4 + i for i in range(1, max_bullets + 1)}
    for row_cells in ws.iter_rows(min_row=2):
        for cell in row_cells:
            if cell.column in wrap_cols:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    logger.info("Built copy XLSX: %d row(s)", len(done))
    return buf.getvalue()
