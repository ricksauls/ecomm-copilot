"""Render a Business Performance Report to a 16:9 PowerPoint deck (python-pptx).

Takes the same dict as :mod:`app.report_pdf` (from
:func:`app.report.assemble_business_report`) and produces a native, fully
editable ``.pptx`` — every title, table, number and chart is a real PowerPoint
object the user can tweak, which is the point of offering PPTX alongside the PDF.

Mirrors the PDF deck slide-for-slide and uses the same DISCOtech palette, so the
two downloads are visually consistent. Pure-Python (python-pptx + lxml, no
Office/COM, no browser), so it runs in a web request or the background worker on
the droplet. Share of shelf is drawn as a native stacked-column chart (editable
in PowerPoint); everything else is shapes, text boxes and tables.
"""

import io
import logging

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Inches, Pt

logger = logging.getLogger(__name__)

# ── 16:9 widescreen geometry (matches the PDF). All layout math is in inches. ─
PAGE_W, PAGE_H = 13.333, 7.5
MARGIN = 0.6
FONT = "Inter"  # self-hosted brand font; PowerPoint falls back cleanly if absent

# ── DISCOtech palette (see app/static/css/tokens.css). ───────────────────────
MIDNIGHT = RGBColor(0x0B, 0x12, 0x20)
ACCENT = RGBColor(0x2E, 0x6D, 0xFF)
INK = RGBColor(0x0B, 0x12, 0x20)
SLATE = RGBColor(0x64, 0x74, 0x8B)
COOL = RGBColor(0x94, 0xA3, 0xB8)
BORDER = RGBColor(0xD7, 0xDE, 0xEA)
DIVIDER = RGBColor(0xEE, 0xF2, 0xF8)
CARD_BG = RGBColor(0xF3, 0xF6, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PALE = RGBColor(0xEA, 0xEF, 0xF7)
GREEN = RGBColor(0x16, 0xA3, 0x4A)
RED = RGBColor(0xFF, 0x3B, 0x30)
AMBER = RGBColor(0xD9, 0x77, 0x06)
GREEN_BG = RGBColor(0xE9, 0xF7, 0xEF)
AMBER_BG = RGBColor(0xFD, 0xF4, 0xE3)

_TONE = {"good": GREEN, "warn": AMBER, "bad": RED, "muted": COOL, "info": ACCENT,
         "up": GREEN, "down": RED, "flat": COOL}
_IMPACT = {"High": RED, "Medium": AMBER, "Low": SLATE}


# ── Low-level slide helpers ──────────────────────────────────────────────────

def _slide(prs):
    """A fresh blank slide (layout 6 is the blank layout in the default template)."""
    return prs.slides.add_slide(prs.slide_layouts[6])


def _fill_bg(slide, color):
    """Paint the whole slide with one color via a full-bleed borderless rectangle."""
    shp = slide.shapes.add_shape(1, 0, 0, Inches(PAGE_W), Inches(PAGE_H))  # 1 = rectangle
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _rect(slide, x, y, w, h, fill=None, line=None, line_w=0.75):
    """A positioned rectangle (card / pill / bar). Dimensions in inches."""
    shp = slide.shapes.add_shape(1, Inches(x), Inches(y), Inches(w), Inches(h))
    shp.shadow.inherit = False
    if fill is None:
        shp.fill.background()
    else:
        shp.fill.solid()
        shp.fill.fore_color.rgb = fill
    if line is None:
        shp.line.fill.background()
    else:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w)
    return shp


def _text(slide, x, y, w, h, text, size, color, *, bold=False, align=PP_ALIGN.LEFT,
          anchor=MSO_ANCHOR.TOP, spacing=None, font=FONT):
    """A single-run text box. Returns the shape for further tweaks."""
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font
    run.font.color.rgb = color
    if spacing is not None:
        _set_tracking(run, spacing)
    return box


def _set_tracking(run, pts):
    """Apply letter-spacing (tracking) to a run — python-pptx has no API for it,
    so set the ``spc`` attribute on the run's ``<a:rPr>`` directly (EMU-ish: the
    attribute is in 1/100 pt)."""
    run.font._rPr.set("spc", str(int(pts * 100)))


def _label(slide, x, y, text, color=COOL, size=8.5, w=4.0):
    """The tracked-uppercase label used above most blocks."""
    _text(slide, x, y, w, 0.25, text.upper(), size, color, bold=True, spacing=1.0)


def _heading(slide, title, sub=None):
    """Accent rule + title (+ optional subtitle); returns the y where body starts."""
    _rect(slide, MARGIN, MARGIN, 0.47, 0.04, fill=ACCENT)
    _text(slide, MARGIN, MARGIN + 0.08, PAGE_W - 2 * MARGIN, 0.6, title, 26, INK, bold=True)
    y = MARGIN + 0.72
    if sub:
        _text(slide, MARGIN, y, PAGE_W - 2 * MARGIN, 0.3, sub, 12, SLATE)
        y += 0.36
    return y + 0.1


def _bullets(slide, x, y, w, h, title, color, bg, bullets):
    """A titled box of bullet lines (Wins / Attention)."""
    _rect(slide, x, y, w, h, fill=bg)
    _text(slide, x + 0.22, y + 0.18, w - 0.44, 0.3, title, 11, color, bold=True)
    box = slide.shapes.add_textbox(Inches(x + 0.22), Inches(y + 0.52),
                                   Inches(w - 0.44), Inches(h - 0.7))
    tf = box.text_frame
    tf.word_wrap = True
    for i, b in enumerate(bullets):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(5)
        run = p.add_run()
        run.text = f"•  {b}"
        run.font.size = Pt(10.5)
        run.font.name = FONT
        run.font.color.rgb = INK


# ── Slides ───────────────────────────────────────────────────────────────────

def _title_slide(prs, meta):
    s = _slide(prs)
    _fill_bg(s, MIDNIGHT)
    _rect(s, MARGIN, MARGIN, 0.75, 0.055, fill=ACCENT)
    _text(s, MARGIN, MARGIN + 0.12, 6, 0.4, "DISCOTECH", 15, PALE, bold=True, spacing=2.0)
    _text(s, MARGIN, PAGE_H / 2 - 0.9, 11, 1.0, "Performance Report", 46, WHITE, bold=True)
    _text(s, MARGIN, PAGE_H / 2 + 0.1, 11, 0.5, "Search & Content Readiness", 20, COOL)
    _rect(s, MARGIN, PAGE_H - MARGIN - 0.9, 0.4, 0.03, fill=ACCENT)
    line = meta["brand_label"]
    if meta.get("period_label"):
        line += f"  ·  {meta['period_label']}"
    _text(s, MARGIN, PAGE_H - MARGIN - 0.75, 11, 0.4, line, 16, WHITE, bold=True)
    _text(s, MARGIN, PAGE_H - MARGIN - 0.38, 11, 0.3, f"Generated {meta['generated_at']}",
          10, SLATE)


def _agenda_slide(prs, has_data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    _heading(s, "What's Inside")
    rows = [("01", "Summary", "Headline KPIs and this period's highlights")]
    if has_data:
        rows.append(("02", "Product Discovery", "Search ranking and share of digital shelf"))
    rows.append((f"{len(rows) + 1:02d}", "Content Readiness",
                 "Primary content quality, scored per product"))
    rows.append((f"{len(rows) + 1:02d}", "Action Plan",
                 "Prioritized recommendations and next steps"))
    y = 2.0
    for num, title, desc in rows:
        _text(s, MARGIN, y, 0.8, 0.5, num, 22, ACCENT, bold=True)
        _text(s, MARGIN + 0.75, y, 9, 0.35, title, 17, INK, bold=True)
        _text(s, MARGIN + 0.75, y + 0.32, 9, 0.3, desc, 11, SLATE)
        y += 0.92


def _kpi_card(s, x, y, w, h, label, card):
    _rect(s, x, y, w, h, fill=CARD_BG, line=BORDER)
    _label(s, x + 0.22, y + 0.2, label, w=w - 0.4)
    value = card.get("value") if card.get("available") else "—"
    _text(s, x + 0.2, y + 0.55, w - 0.4, 0.8, value or "—", 40, INK if card.get("available") else COOL, bold=True)
    delta = card.get("delta") or {}
    if delta.get("display"):
        _text(s, x + 0.22, y + h - 0.72, w - 0.4, 0.3, delta["display"], 10,
              _TONE.get(delta.get("dir"), COOL), bold=True)
    _text(s, x + 0.22, y + h - 0.42, w - 0.4, 0.3, card.get("caption", ""), 9, SLATE)


def _summary_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    _heading(s, "Summary", "Overall KPI health")
    cards = data["cards"]
    usable = PAGE_W - 2 * MARGIN
    group_gap, gap, card_h = 0.5, 0.25, 2.0
    disc_w = (usable - group_gap) * 0.62
    conv_w = (usable - group_gap) * 0.38
    top = 2.0
    _label(s, MARGIN, top - 0.28, "Discovery Drivers", color=ACCENT, size=9)
    _label(s, MARGIN + disc_w + group_gap, top - 0.28, "Conversion Drivers", color=ACCENT, size=9)
    cw = (disc_w - gap) / 2
    _kpi_card(s, MARGIN, top, cw, card_h, "Avg Search Rank", cards["avg_rank"])
    _kpi_card(s, MARGIN + cw + gap, top, cw, card_h, "Share of Page 1", cards["shelf_share"])
    _kpi_card(s, MARGIN + disc_w + group_gap, top, conv_w, card_h,
              "Content Readiness", cards["content_readiness"])


def _highlights_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Highlights")
    hl = data["highlights"]
    usable = PAGE_W - 2 * MARGIN
    total_h = PAGE_H - y - MARGIN
    box_h = (total_h - 0.25) / 2
    _bullets(s, MARGIN, y, usable, box_h, "WINS", GREEN, GREEN_BG, hl["wins"])
    _bullets(s, MARGIN, y + box_h + 0.25, usable, box_h, "ATTENTION NEEDED",
             AMBER, AMBER_BG, hl["attention"])


def _section_slide(prs, num, title, desc):
    s = _slide(prs)
    _fill_bg(s, MIDNIGHT)
    _text(s, MARGIN, PAGE_H / 2 - 0.95, 8, 0.3, f"SECTION {num}", 12, ACCENT, bold=True, spacing=2.0)
    _text(s, MARGIN, PAGE_H / 2 - 0.55, 11, 0.8, title, 40, WHITE, bold=True)
    if desc:
        _text(s, MARGIN, PAGE_H / 2 + 0.3, 11, 0.4, desc, 13, COOL)


def _table(slide, x, y, w, header, rows, col_fracs, *, max_rows=12):
    """A compact data table, top-anchored at (x, y)."""
    rows = rows[:max_rows]
    nrows = len(rows) + 1
    row_h = 0.3
    gf = slide.shapes.add_table(nrows, len(header), Inches(x), Inches(y),
                                Inches(w), Inches(row_h * nrows))
    tbl = gf.table
    tbl.first_row = False
    tbl.horz_banding = False
    for ci, frac in enumerate(col_fracs):
        tbl.columns[ci].width = Inches(w * frac)
    for ci, htext in enumerate(header):
        cell = tbl.cell(0, ci)
        _fill_cell(cell, CARD_BG)
        _cell_text(cell, htext, 8.5, SLATE, bold=True)
    for ri, row in enumerate(rows, start=1):
        for ci, val in enumerate(row):
            cell = tbl.cell(ri, ci)
            _fill_cell(cell, WHITE)
            _cell_text(cell, str(val), 9, INK if ci == 0 else SLATE)
    return gf


def _fill_cell(cell, color):
    cell.fill.solid()
    cell.fill.fore_color.rgb = color
    cell.margin_left = cell.margin_right = Inches(0.08)
    cell.margin_top = cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE


def _cell_text(cell, text, size, color, *, bold=False):
    cell.text = ""
    p = cell.text_frame.paragraphs[0]
    run = p.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = FONT
    run.font.color.rgb = color


def _rank_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Search Rank", "Average page-1 position across tracked keywords")
    disc = data["discovery"]
    rows = [[r["brand_name"], r.get("type", ""),
             f"#{r['avg_position']}" if r.get("avg_position") is not None else "—"]
            for r in disc["avg_ranks"]]
    _table(s, MARGIN, y, PAGE_W - 2 * MARGIN, ["Brand", "Type", "Avg rank"], rows,
           [0.5, 0.25, 0.25], max_rows=11)


def _shelf_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Share of Digital Shelf", "Your page-1 slots vs. the whole shelf")
    disc = data["discovery"]
    usable = PAGE_W - 2 * MARGIN
    left_w = usable * 0.52
    rows = [[r["brand_name"], f"{r['organic_share']}%", f"{r['sponsored_share']}%",
             f"{r['total_share']}%"] for r in disc["sos_rows"]]
    _table(s, MARGIN, y, left_w, ["Brand", "Organic", "Sponsored", "Total"], rows,
           [0.4, 0.2, 0.2, 0.2], max_rows=10)
    # Native stacked-column chart of organic + sponsored share per brand.
    sos = disc["sos_rows"]
    if sos:
        cd = CategoryChartData()
        cd.categories = [r["brand_name"] for r in sos]
        cd.add_series("Organic", tuple(r["organic_share"] for r in sos))
        cd.add_series("Sponsored", tuple(r["sponsored_share"] for r in sos))
        cx = MARGIN + left_w + 0.3
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_STACKED, Inches(cx), Inches(y),
                                Inches(usable - left_w - 0.3), Inches(3.6), cd)
        chart = gf.chart
        chart.has_title = False
        chart.has_legend = True
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.plots[0].series[0].format.fill.solid()
        chart.plots[0].series[0].format.fill.fore_color.rgb = INK
        chart.plots[0].series[1].format.fill.solid()
        chart.plots[0].series[1].format.fill.fore_color.rgb = COOL


def _mix_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Organic vs. Paid Placement Mix")
    mix = data["discovery"]["mix"]
    usable = PAGE_W - 2 * MARGIN
    cw, ch = (usable - 0.3) / 2, 2.0
    specs = [("Organic Placements", mix["organic_pct"], mix["organic_share"],
              "Primary discovery channel"),
             ("Sponsored Placements", mix["sponsored_pct"], mix["sponsored_share"],
              "Paid search results")]
    for i, (label, pct, share, note) in enumerate(specs):
        x = MARGIN + i * (cw + 0.3)
        _rect(s, x, y, cw, ch, fill=CARD_BG, line=BORDER)
        _label(s, x + 0.22, y + 0.2, label, w=cw - 0.4)
        _text(s, x + 0.2, y + 0.55, cw - 0.4, 0.9, f"{pct}%", 44, INK, bold=True)
        _text(s, x + 0.22, y + ch - 0.7, cw - 0.4, 0.3, note, 10, SLATE)
        _text(s, x + 0.22, y + ch - 0.42, cw - 0.4, 0.3, f"{share}% share of page 1",
              10, ACCENT, bold=True)


def _callout(slide, x, y, w, tone, headline, body):
    """An insight callout (accent bar + headline + body). Returns height used (in)."""
    color = _TONE.get(tone, ACCENT)
    _rect(slide, x, y, 0.05, 0.95, fill=color)
    _text(slide, x + 0.18, y, w - 0.3, 0.3, headline, 11.5, color, bold=True)
    _text(slide, x + 0.18, y + 0.3, w - 0.3, 0.7, body, 10, INK)
    return 1.1


def _insights_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Discovery Key Insights")
    usable = PAGE_W - 2 * MARGIN
    for ins in data["discovery"]["insights"]:
        used = _callout(s, MARGIN, y, usable, ins["tone"], ins["headline"], ins["body"])
        y += used + 0.1
        if y > PAGE_H - MARGIN:
            break


def _readiness_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    r = data["readiness"]
    sub = (f"Overall {r['overall']}/100 across {r['scored_count']} product"
           f"{'' if r['scored_count'] == 1 else 's'}" if r["available"]
           else "No products scored yet")
    y = _heading(s, "Content Readiness Scorecard", sub)
    _label(s, MARGIN, y, "Primary Content")
    ry = y + 0.3
    usable = PAGE_W - 2 * MARGIN
    bar_x = MARGIN + 2.0
    bar_w = usable - 2.0 - 2.0
    for d in r["dimensions"]:
        _text(s, MARGIN, ry, 1.9, 0.3, d["label"], 11, INK, bold=True)
        _rect(s, bar_x, ry + 0.05, bar_w, 0.12, fill=DIVIDER)
        if d["avg"]:
            _rect(s, bar_x, ry + 0.05, bar_w * min(d["avg"], 100) / 100, 0.12,
                  fill=_TONE.get(d["grade"]["tone"], COOL))
        _text(s, bar_x + bar_w + 0.15, ry, 0.5, 0.3,
              f"{d['avg']}" if d["avg"] is not None else "n/a", 12, INK, bold=True)
        _text(s, bar_x + bar_w + 0.7, ry, 1.3, 0.3, d["grade"]["label"], 9,
              _TONE.get(d["grade"]["tone"], COOL), bold=True)
        ry += 0.5
    if not r["dimensions"]:
        _text(s, MARGIN, ry, usable, 0.3,
              "Score this brand's products to populate the scorecard.", 11, SLATE)


def _item_scores_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Content Scores by Item")
    rows = [[it["title"] or it["item_id"], f"{it['score']}", it["grade"]["label"]]
            for it in data["readiness"]["items"]]
    _table(s, MARGIN, y, PAGE_W - 2 * MARGIN, ["Item", "Score", "Grade"], rows,
           [0.64, 0.16, 0.20], max_rows=12)


def _fixes_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Listing Readiness Priority Fixes")
    fixes = data["priority_fixes"]
    if fixes:
        rows = [[f["item"], f["issue"], f["impact"], f["action"]] for f in fixes]
        _table(s, MARGIN, y, PAGE_W - 2 * MARGIN,
               ["Item", "Issue", "Impact", "Action required"], rows,
               [0.24, 0.28, 0.12, 0.36], max_rows=12)
    else:
        _text(s, MARGIN, y + 0.3, PAGE_W - 2 * MARGIN, 0.4,
              "No content gaps flagged — all scored products meet target.", 12, SLATE)


def _recs_slide(prs, data):
    s = _slide(prs)
    _fill_bg(s, WHITE)
    y = _heading(s, "Strategic Recommendations")
    usable = PAGE_W - 2 * MARGIN
    for rec in data["recommendations"]:
        color = _IMPACT.get(rec["priority"], SLATE)
        _rect(s, MARGIN, y, 0.75, 0.22, fill=color)
        _text(s, MARGIN + 0.1, y + 0.02, 0.6, 0.2, rec["priority"].upper(), 8, WHITE, bold=True)
        _text(s, MARGIN + 0.9, y - 0.02, usable - 0.9, 0.3, rec["title"], 13, INK, bold=True)
        _text(s, MARGIN + 0.9, y + 0.28, usable - 0.9, 0.7, rec["body"], 10, SLATE)
        y += 1.05
        if y > PAGE_H - MARGIN:
            break


def _close_slide(prs, meta):
    s = _slide(prs)
    _fill_bg(s, MIDNIGHT)
    _text(s, MARGIN, PAGE_H / 2 - 0.5, 11, 0.9, "Questions?", 44, WHITE, bold=True)
    _text(s, MARGIN, PAGE_H / 2 + 0.4, 11, 0.4,
          f"{meta['brand_label']} · {meta.get('period_label') or 'Current'} · DISCOtech",
          13, COOL)


def build_business_report_pptx(data: dict) -> bytes:
    """Render the assembled report dict to a PPTX deck; return the bytes.

    Same adaptive slide set as the PDF: the Product Discovery section appears only
    when CI monitoring has a completed period. Every slide handles empty inputs,
    so a thin-data brand still produces a complete, editable deck.
    """
    meta = data["meta"]
    has_data = meta["has_data"]
    prs = Presentation()
    prs.slide_width = Inches(PAGE_W)
    prs.slide_height = Inches(PAGE_H)

    _title_slide(prs, meta)
    _agenda_slide(prs, has_data)
    _summary_slide(prs, data)
    _highlights_slide(prs, data)
    sec = 1
    if has_data:
        _section_slide(prs, f"{sec:02d}", "Product Discovery",
                       "Search ranking and share of digital shelf")
        _rank_slide(prs, data)
        _shelf_slide(prs, data)
        _mix_slide(prs, data)
        _insights_slide(prs, data)
        sec += 1
    _section_slide(prs, f"{sec:02d}", "Content Readiness",
                   "Primary content quality, scored per product")
    _readiness_slide(prs, data)
    _item_scores_slide(prs, data)
    _fixes_slide(prs, data)
    sec += 1
    _section_slide(prs, f"{sec:02d}", "Action Plan",
                   "Prioritized recommendations and next steps")
    _recs_slide(prs, data)
    _close_slide(prs, meta)

    buffer = io.BytesIO()
    prs.save(buffer)
    logger.info("Built business report PPTX: brand=%s period=%s slides=%d",
                meta["brand_label"], meta["period"], len(prs.slides._sldIdLst))
    return buffer.getvalue()
