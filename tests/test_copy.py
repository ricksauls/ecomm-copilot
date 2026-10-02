"""Route tests for PDP Copy Content Creation (intake, results, generate, cross-link)."""

import json

from app import copy_jobs, jobs
from app.db import get_db


def test_copy_routes_require_login(client):
    assert client.get("/app/pdp-copy").status_code == 302
    assert client.get("/app/pdp-copy/results").status_code == 302
    assert client.get("/app/pdp-copy/results.pdf").status_code == 302
    assert client.get("/app/pdp-copy/results.xlsx").status_code == 302
    assert client.get("/app/pdp-copy/results.csv").status_code == 302
    assert client.post("/app/pdp-copy/generate").status_code == 302
    assert client.post("/app/pdp-scoring/create-copy").status_code == 302


def test_copy_intake_renders(client, auth):
    auth.register()
    resp = client.get("/app/pdp-copy")
    assert resp.status_code == 200
    assert b"Create Product Detail Page Copy Content" in resp.data
    # The fetch button carries the exact requested label.
    assert b"Get Current Copy Content" in resp.data
    # CSV cap mirrors app.pdp.MAX_ITEMS (100).
    assert b"up to 100" in resp.data
    # The copy intake keeps the optional brand field.
    assert b'name="brand"' in resp.data


def test_copy_intake_persists_entered_brand(client, auth):
    auth.register()
    client.post(
        "/app/pdp-copy",
        data={"urls": "https://www.walmart.com/ip/10294528", "brand": "  Tabasco "},
    )
    with client.application.app_context():
        row = get_db().execute(
            "SELECT brand FROM copy_items WHERE item_id = '10294528'"
        ).fetchone()
        assert row["brand"] == "Tabasco"  # trimmed by clean_brand


def test_copy_intake_enqueues_and_redirects(client, auth):
    auth.register()
    resp = client.post(
        "/app/pdp-copy",
        data={"urls": "https://www.walmart.com/ip/10294528"},
    )
    # Redirects to the results page on success.
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/app/pdp-copy/results")
    # The results page then shows the queued item.
    results = client.get("/app/pdp-copy/results")
    assert results.status_code == 200
    assert b"Fetching" in results.data


def test_copy_intake_rejects_empty(client, auth):
    auth.register()
    resp = client.post("/app/pdp-copy", data={"urls": "not-a-url"})
    assert resp.status_code == 400
    assert b"No valid item URLs" in resp.data


def test_copy_status_json(client, auth):
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    resp = client.get("/app/pdp-copy/status")
    data = resp.get_json()
    assert data["pending"] is True
    assert len(data["items"]) == 1
    assert data["items"][0]["status"] == "queued"


def test_generate_advances_fetched_items(client, auth, app):
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    # Simulate the worker having fetched the current copy.
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT id FROM copy_items").fetchone()
        copy_jobs.save_current_copy(
            db, row["id"], title="P", current={"record": {"url": "u"}},
            current_overall=50, keywords=[], next_status="fetched",
        )
    resp = client.post("/app/pdp-copy/generate")
    assert resp.status_code == 302
    with app.app_context():
        status = get_db().execute("SELECT status FROM copy_items").fetchone()["status"]
        assert status == "gen_queued"


def test_copy_results_pdf_download(client, auth, app):
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    # Simulate a completed item (current + new copy + scores).
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT id FROM copy_items ORDER BY id DESC LIMIT 1").fetchone()
        copy_jobs.save_current_copy(
            db, row["id"], title="Acme Widget",
            current={"title": "OLD", "bullets": ["a"], "description": "old desc",
                     "record": {"url": "u"}},
            current_overall=60, keywords=[], next_status="fetched",
        )
        copy_jobs.save_generated_copy(
            db, row["id"],
            new={"title": "NEW", "bullets": ["b1", "b2"], "description": "new desc"},
            projected_overall=88,
        )
    resp = client.get("/app/pdp-copy/results.pdf")
    assert resp.status_code == 200
    assert resp.mimetype == "application/pdf"
    assert resp.data[:5] == b"%PDF-"  # real PDF, not an error page
    assert "attachment" in resp.headers.get("Content-Disposition", "")


def test_copy_results_shows_download_only_when_done(client, auth, app):
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    # Still fetching -> no Download PDF button yet.
    assert b"Download PDF" not in client.get("/app/pdp-copy/results").data
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT id FROM copy_items ORDER BY id DESC LIMIT 1").fetchone()
        copy_jobs.save_current_copy(
            db, row["id"], title="P",
            current={"title": "OLD", "bullets": ["a"], "description": "d",
                     "record": {"url": "u"}},
            current_overall=60, keywords=[], next_status="fetched",
        )
        copy_jobs.save_generated_copy(
            db, row["id"], new={"title": "NEW", "bullets": ["b"], "description": "d"},
            projected_overall=88,
        )
    # Now done -> the button appears.
    assert b"Download PDF" in client.get("/app/pdp-copy/results").data


def test_scoring_cross_link_without_record_falls_back_to_fetch(client, auth, app):
    # A scored item with no stored record (older score) takes the fetch path:
    # the copy row starts 'queued' and will re-fetch the PDP.
    auth.register()
    with app.app_context():
        db = get_db()
        uid = db.execute("SELECT id FROM users").fetchone()["id"]
        ids = jobs.enqueue_items(db, uid, [{"url": "https://www.walmart.com/ip/5", "item": "5"}])
        jobs.save_result(db, ids[0], 70, {"overall": 70, "dimensions": []}, "Prod")
        scored_id = ids[0]

    resp = client.post("/app/pdp-scoring/create-copy", data={"item_ids": str(scored_id)})
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/app/pdp-copy/results")
    with app.app_context():
        row = get_db().execute("SELECT url, auto_generate, status FROM copy_items").fetchone()
        assert row["url"] == "https://www.walmart.com/ip/5"
        assert row["auto_generate"] == 1
        assert row["status"] == "queued"  # fetch path (no record to reuse)


def test_scoring_cross_link_reuses_stored_record_without_refetch(client, auth, app):
    # A scored item that carries the fetched record goes straight to generation —
    # no second PDP fetch — with its current copy pre-populated from the score.
    auth.register()
    with app.app_context():
        db = get_db()
        uid = db.execute("SELECT id FROM users").fetchone()["id"]
        ids = jobs.enqueue_items(db, uid, [{"url": "https://www.walmart.com/ip/5", "item": "5"}])
        record = {
            "title": "Prod", "bullets": ["Bullet one"], "description": "Desc text",
            "target_keywords": ["kw1", "kw2"],
        }
        jobs.save_result(db, ids[0], 70, {"overall": 70, "dimensions": []}, "Prod",
                         record=record)
        scored_id = ids[0]

    resp = client.post("/app/pdp-scoring/create-copy", data={"item_ids": str(scored_id)})
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/app/pdp-copy/results")
    with app.app_context():
        row = get_db().execute(
            "SELECT status, auto_generate, current_json, current_overall, keywords_json "
            "FROM copy_items"
        ).fetchone()
        assert row["status"] == "gen_queued"   # skips the fetch phase
        assert row["auto_generate"] == 1
        assert row["current_overall"] == 70
        current = json.loads(row["current_json"])
        assert current["title"] == "Prod"
        assert current["bullets"] == ["Bullet one"]
        assert current["description"] == "Desc text"
        assert current["record"]["target_keywords"] == ["kw1", "kw2"]
        assert current["captured_at"]  # capture time recorded for the UI note
        assert json.loads(row["keywords_json"]) == ["kw1", "kw2"]


def test_scoring_cross_link_ignores_other_users_items(client, auth, app):
    auth.register()
    with app.app_context():
        from app.users import create_local_user
        other = create_local_user("someone@else.com", "password123")
        db = get_db()
        ids = jobs.enqueue_items(db, other, [{"url": "https://www.walmart.com/ip/9", "item": "9"}])
        jobs.save_result(db, ids[0], 70, {"overall": 70, "dimensions": []}, "Prod")
        other_id = ids[0]

    # Trying to copy another user's item creates nothing (IDOR guard) and bounces
    # back to the scoring results.
    resp = client.post("/app/pdp-scoring/create-copy", data={"item_ids": str(other_id)})
    assert resp.status_code == 302
    assert resp.headers["Location"].endswith("/app/pdp-scoring/results")
    with app.app_context():
        assert get_db().execute("SELECT COUNT(*) AS c FROM copy_items").fetchone()["c"] == 0


def test_fmt_captured_formats_and_tolerates_bad_input():
    from app.routes.pages import _fmt_captured

    assert _fmt_captured("2026-10-01 16:12:00") == "Oct 01, 2026 04:12 PM UTC"
    assert _fmt_captured(None) is None
    assert _fmt_captured("not a date") == "not a date"  # never 500s the page


def test_copy_results_shows_progress_while_pending(client, auth):
    # A freshly enqueued (unfetched) copy batch is pending: the subtitle shows the
    # in-progress text, a progress bar (0 of 1 so far) with a time estimate, and no
    # flashing class (parity with the Content Scores screen).
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    resp = client.get("/app/pdp-copy/results")
    assert b"about 1 minute per 5 items" in resp.data
    assert b'role="progressbar"' in resp.data
    assert b"0 of 1 complete" in resp.data
    assert b"less than a minute left" in resp.data
    assert b"subtitle-scoring" not in resp.data  # no text flashing on this screen


def test_copy_export_buttons_appear_only_when_done(client, auth, app):
    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/1"})
    before = client.get("/app/pdp-copy/results").data
    assert b"Download To Excel" not in before
    assert b"Download CSV" not in before
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT id FROM copy_items ORDER BY id DESC LIMIT 1").fetchone()
        copy_jobs.save_current_copy(
            db, row["id"], title="P",
            current={"title": "O", "bullets": ["a"], "description": "d", "record": {"url": "u"}},
            current_overall=60, keywords=[], next_status="fetched",
        )
        copy_jobs.save_generated_copy(
            db, row["id"], new={"title": "N", "bullets": ["b"], "description": "d"},
            projected_overall=88,
        )
    after = client.get("/app/pdp-copy/results").data
    assert b"Download To Excel" in after
    assert b"Download CSV" in after


def test_copy_exports_contain_new_copy(client, auth, app):
    import io

    from openpyxl import load_workbook

    auth.register()
    client.post("/app/pdp-copy", data={"urls": "https://www.walmart.com/ip/10294528"})
    with app.app_context():
        db = get_db()
        row = db.execute("SELECT id FROM copy_items ORDER BY id DESC LIMIT 1").fetchone()
        copy_jobs.save_current_copy(
            db, row["id"], title="Acme Widget",
            current={"title": "OLD", "bullets": ["a"], "description": "old",
                     "record": {"url": "u"}},
            current_overall=60, keywords=[], next_status="fetched",
        )
        copy_jobs.save_generated_copy(
            db, row["id"],
            new={"title": "Acme Widget Deluxe",
                 "bullets": ["Durable steel", "Easy setup"],
                 "description": "A great widget."},
            projected_overall=88,
        )

    # Excel: a real .xlsx (zip) with the new copy mapped to Walmart content columns.
    xlsx = client.get("/app/pdp-copy/results.xlsx")
    assert xlsx.status_code == 200
    assert xlsx.mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert xlsx.data[:2] == b"PK"
    assert "attachment" in xlsx.headers.get("Content-Disposition", "")
    ws = load_workbook(io.BytesIO(xlsx.data)).active
    header = [c.value for c in ws[1]]
    assert header[:4] == ["Item ID", "Product URL", "Product Name", "Site Description"]
    assert "Key Feature 1" in header and "Key Feature 2" in header
    row2 = [c.value for c in ws[2]]
    assert row2[0] == "10294528"
    assert row2[2] == "Acme Widget Deluxe"
    assert row2[3] == "A great widget."
    assert "Durable steel" in row2 and "Easy setup" in row2

    # CSV: same content, no Excel required.
    csv_resp = client.get("/app/pdp-copy/results.csv")
    assert csv_resp.status_code == 200
    assert csv_resp.mimetype == "text/csv"
    body = csv_resp.data.decode("utf-8-sig")
    assert "Product Name" in body
    assert "Acme Widget Deluxe" in body
    assert "A great widget." in body
    assert "Durable steel" in body
